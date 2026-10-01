import datetime as dt
import io
import unittest
import uuid
import zipfile

import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

import export


class ExportTextAndPhotoTests(unittest.TestCase):
    def state(self):
        return {
            "df_weighings": pd.DataFrame([{
                "__weighing_id": uuid.uuid4().hex,
                "N° échantillon": "1", "Classe de matériau": "Paper",
                "Contenant utilisé": "", "Poids brut": 1.0,
                "Tare": 0.0, "Poids net": 1.0,
            }]),
            "df_collect_times": pd.DataFrame(columns=["Echantillon"]),
            "saved_sensor_name": "Sensor", "saved_test_date": dt.date(2026, 10, 1),
            "saved_operator_name": "Alice", "saved_nb_sample": 1,
            "_skip_collect_times_value": True, "saved_workflow": 0,
            "saved_workflow_order": 0, "material_classes": ["Paper"],
            "facility_name": "Site", "user_timezone": "Europe/Zurich",
        }

    def photo(self, format):
        buffer = io.BytesIO()
        Image.new("RGB", (3, 3), color="green").save(buffer, format=format)
        return buffer.getvalue()

    def test_unicode_survives_in_pdf_metadata_material_and_comment(self):
        state = self.state()
        state["saved_operator_name"] = "Łukasz García"
        state["facility_name"] = "Zürich – Genève"
        state["saved_sensor_name"] = "Capteur № 1"
        state["df_weighings"]["Classe de matériau"] = "Plastique ♻"
        state["_global_comment_value"] = "Collecte à 10 h — déjà triée."
        with zipfile.ZipFile(export.build_zip_export("unicode", state=state)) as archive:
            reader = PdfReader(io.BytesIO(archive.read("unicode.pdf")))
            text = "\n".join(page.extract_text() for page in reader.pages)
        for expected in (state["saved_operator_name"], state["facility_name"],
                         state["saved_sensor_name"], "Plastique ♻", state["_global_comment_value"]):
            self.assertIn(expected, text)

    def test_unsupported_glyph_does_not_block_zip_and_excel_keeps_original(self):
        state = self.state()
        state["_global_comment_value"] = "Rare symbol: \U0001fae8"
        with zipfile.ZipFile(export.build_zip_export("glyph", state=state)) as archive:
            reader = PdfReader(io.BytesIO(archive.read("glyph.pdf")))
            text = "\n".join(page.extract_text() for page in reader.pages)
            self.assertIn("Rare symbol: ?", text)
            worksheet = load_workbook(io.BytesIO(archive.read("glyph.xlsx")))["Metadata"]
            metadata = dict(worksheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(metadata["Global comment"], state["_global_comment_value"])

    def test_excel_free_text_stays_literal_and_weights_stay_numeric(self):
        for value in ("=1+1", '=HYPERLINK("https://example.com","text")',
                      "+1+1", "-1+1", "@SUM(A1)", "#N/A", "Łukasz ♻"):
            with self.subTest(value=value):
                state = self.state()
                for key in ("saved_operator_name", "saved_sensor_name", "facility_name", "_global_comment_value"):
                    state[key] = value
                state["df_weighings"]["Classe de matériau"] = value
                workbook = load_workbook(export.build_excel_export(state=state))
                for sheet in workbook:
                    for row in sheet:
                        for cell in row:
                            self.assertNotIn(cell.data_type, ("f", "e"))
                global_sheet = workbook["Global results"]
                self.assertEqual(global_sheet["F2"].value, value)
                self.assertEqual(global_sheet["F2"].data_type, "s")
                self.assertEqual(global_sheet["G2"].data_type, "n")
                metadata = dict(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
                for key in ("Facility name", "Operator name", "Sensor name", "Global comment"):
                    self.assertEqual(metadata[key], value)
                self.assertIn(value, [cell.value for row in workbook["Sample 1"] for cell in row])

    def test_same_material_photos_are_all_preserved_with_detected_formats(self):
        state = self.state()
        jpeg, png = self.photo("JPEG"), self.photo("PNG")
        first = state["df_weighings"].iloc[0].to_dict()
        first["Image"] = jpeg
        second = {**first, "__weighing_id": uuid.uuid4().hex, "Image": png}
        state["df_weighings"] = pd.DataFrame([first, second])
        original = state["df_weighings"].copy(deep=True)
        with zipfile.ZipFile(export.build_zip_export("photos", state=state)) as archive:
            photos = [name for name in archive.namelist() if name.startswith("images/")]
            self.assertEqual(len(photos), 2)
            self.assertEqual(archive.read(f"images/Paper_{first['__weighing_id']}.jpg"), jpeg)
            self.assertEqual(archive.read(f"images/Paper_{second['__weighing_id']}.png"), png)
        pd.testing.assert_frame_equal(state["df_weighings"], original)

    def test_photo_names_cannot_create_paths_and_unknown_bytes_are_preserved(self):
        state = self.state()
        state["df_weighings"]["Classe de matériau"] = '../../folder\\material:<>?*'
        state["df_weighings"]["Image"] = [b"unrecognised original attachment"]
        with zipfile.ZipFile(export.build_zip_export("../unsafe\\report", state=state)) as archive:
            self.assertEqual(len(archive.namelist()), 3)
            for name in archive.namelist():
                self.assertNotIn("\\", name)
                self.assertNotIn("..", name.split("/"))
            photos = [name for name in archive.namelist() if name.startswith("images/")]
            self.assertEqual(len(photos), 1)
            self.assertEqual(photos[0].count("/"), 1)
            self.assertTrue(photos[0].endswith(".bin"))
            self.assertEqual(archive.read(photos[0]), b"unrecognised original attachment")


if __name__ == "__main__":
    unittest.main()
